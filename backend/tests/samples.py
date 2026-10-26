"""Synthetic (invented) clinical text and file builders for tests. No real patient data."""
import io

SYNTHETIC_NOTE = (
    "Patient Name: Ramesh Kumar  Age: 54  MRN: 00456712\n"
    "Contact: +91 98765 43210, email ramesh.k@example.com\n"
    "Address: H.No. 12-3-45, Ramnagar Colony, Hanamkonda\n"
    "Seen by Dr. Anita Rao on 12/03/2024.\n"
    "Patient has type 2 diabetes mellitus and hypertension. Started on Metformin 500 mg "
    "twice daily. HbA1c 8.2 %. BP 130/85 mmHg.\n"
)


def make_pdf(text: str = SYNTHETIC_NOTE) -> bytes:
    import pymupdf
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_textbox(pymupdf.Rect(50, 50, 550, 780), text, fontsize=11)
    data = doc.tobytes()
    doc.close()
    return data


def make_blank_pdf() -> bytes:
    import pymupdf
    doc = pymupdf.open()
    doc.new_page()
    data = doc.tobytes()
    doc.close()
    return data


def make_docx(paragraphs=("Discharge summary", "Patient has asthma."),
              table=(("Test", "Result"), ("Hemoglobin", "11.2 g/dL"))) -> bytes:
    from docx import Document
    doc = Document()
    doc.add_paragraph(paragraphs[0])
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, val in enumerate(row):
                t.cell(r, c).text = val
    for p in paragraphs[1:]:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# Synthetic note for NLP tests (all names, numbers and findings are invented).
NLP_NOTE = (
    "Patient Name: Ramesh Kumar  Age: 54  MRN: 00456712\n"
    "Contact: +91 98765 43210\n"
    "Seen by Dr. Anita Rao on 12/03/2024.\n"
    "A 54-year-old male with type 2 diabetes mellitus presented with polyuria and fatigue.\n"
    "He was started on Metformin 500 mg twice daily for type 2 diabetes mellitus.\n"
    "HbA1c 8.2 % consistent with poor glycemic control. "
    "Ultrasound of the abdomen showed fatty liver.\n"
)
